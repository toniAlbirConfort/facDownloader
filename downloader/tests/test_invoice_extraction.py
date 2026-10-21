from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

import fitz
from openpyxl import load_workbook
from PIL import Image
from django.test import SimpleTestCase

from downloader.invoice_extraction import build_invoice_workbook, extract_invoice_fields, parse_invoice_text


INVOICE_TEXT = """Factura de electricidad
Titular: Albir Confort SL
NIF: B12345678
Dirección de suministro: Calle Mayor 10, Benidorm
Número de factura: ELEC-2026-0001
Fecha de emisión: 12/10/2026
Periodo desde: 01/10/2026
Periodo hasta: 31/10/2026
CUPS ES0021000000000000AA
Potencia contratada 4,60 kW
Consumo total 250,50 kWh
Tarifa 2.0TD
Base imponible: 100,00
Impuestos: 21,00
Total a pagar: 121,00 EUR
Vencimiento: 30/10/2026
"""


def write_pdf(path: Path, text: str | None = None, image_only: bool = False):
    pdf = fitz.open()
    page = pdf.new_page()
    if image_only:
        image = Image.new("RGB", (900, 1200), "white")
        from PIL import ImageDraw
        ImageDraw.Draw(image).text((35, 35), text or "", fill="black", spacing=8)
        output = BytesIO()
        image.save(output, format="PNG")
        page.insert_image(page.rect, stream=output.getvalue())
    elif text:
        page.insert_textbox(fitz.Rect(40, 40, 550, 790), text, fontsize=10)
    pdf.save(path)
    pdf.close()


class InvoiceExtractionTests(SimpleTestCase):
    def test_text_pdf_extracts_common_and_electricity_fields_without_ocr(self):
        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "factura.pdf"
            write_pdf(path, INVOICE_TEXT)
            with patch("downloader.invoice_extraction.pytesseract.image_to_string") as ocr:
                fields = extract_invoice_fields(path, "luz", "repsol")

        self.assertEqual(fields["invoice_number"], "ELEC-2026-0001")
        self.assertEqual(fields["issue_date"], "12/10/2026")
        self.assertEqual(fields["total_amount"], "121,00")
        self.assertEqual(fields["holder"], "Albir Confort SL")
        self.assertEqual(fields["consumption_kwh"], "250,50")
        self.assertEqual(fields["extraction_status"], "Correcta")
        ocr.assert_not_called()

    def test_scanned_pdf_uses_local_ocr_fallback(self):
        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "escaneada.pdf"
            write_pdf(path, image_only=True)
            with patch("downloader.invoice_extraction.pytesseract.image_to_string", return_value=INVOICE_TEXT) as ocr:
                fields = extract_invoice_fields(path, "luz", "repsol")

        self.assertGreater(ocr.call_count, 0)
        self.assertEqual(fields["invoice_number"], "ELEC-2026-0001")
        self.assertEqual(fields["extraction_status"], "Correcta")

    def test_missing_or_unreadable_fields_are_marked_for_review(self):
        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "invalid.pdf"
            path.write_bytes(b"not a PDF")
            fields = extract_invoice_fields(path, "agua", "veolia")

        self.assertEqual(fields["extraction_status"], "Revisar")
        self.assertIn("número de factura", fields["missing_fields"])
        self.assertIn("ocr_note", fields)

    def test_parser_includes_service_specific_columns_and_blank_unknowns(self):
        fields = parse_invoice_text("Número de factura: GAS-1\nTotal a pagar: 55,20", "gas", "repsol")

        self.assertIn("cups", fields)
        self.assertIn("meter_number", fields)
        self.assertIn("consumption_kwh", fields)
        self.assertEqual(fields["cups"], "")
        self.assertEqual(fields["extraction_status"], "Revisar")

    def test_workbook_has_one_row_per_pdf_and_review_columns(self):
        with TemporaryDirectory() as temporary_root:
            root = Path(temporary_root)
            path = root / "factura.pdf"
            write_pdf(path, INVOICE_TEXT)
            destination = root / "facturas-luz.xlsx"
            documents = [SimpleNamespace(file_path=str(path), provider="repsol")]
            build_invoice_workbook(documents, "luz", destination)
            workbook = load_workbook(destination, read_only=True)
            sheet = workbook["Luz"]

            self.assertEqual(sheet.max_row, 2)
            headers = [cell.value for cell in sheet[1]]
            values = dict(zip(headers, [cell.value for cell in sheet[2]]))
            self.assertEqual(values["Número de factura"], "ELEC-2026-0001")
            self.assertEqual(values["Total (€)"], 121.0)
            self.assertEqual(values["Estado de extracción"], "Correcta")
            self.assertIn("Potencia contratada (kW)", headers)
            workbook.close()
