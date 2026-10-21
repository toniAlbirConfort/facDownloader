"""Local PDF text/OCR extraction and spreadsheet generation for invoices."""

import logging
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from time import perf_counter

import fitz
import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image

logger = logging.getLogger(__name__)


COMMON_FIELDS = (
    ("Tipo", "invoice_type"),
    ("Proveedor", "provider"),
    ("Titular", "holder"),
    ("NIF/CIF", "tax_id"),
    ("Dirección de suministro", "supply_address"),
    ("Número de factura", "invoice_number"),
    ("Número de contrato", "contract_number"),
    ("Contrato de acceso", "access_contract_number"),
    ("Fecha de emisión", "issue_date"),
    ("Periodo desde", "period_start"),
    ("Periodo hasta", "period_end"),
    ("Base imponible (€)", "tax_base"),
    ("Impuestos (€)", "tax_amount"),
    ("Total (€)", "total_amount"),
    ("Vencimiento", "due_date"),
)

SERVICE_FIELDS = {
    "agua": (
        ("Contador", "meter_number"),
        ("Calibre del contador (mm)", "meter_size_mm"),
        ("Fecha de lectura anterior", "previous_reading_date"),
        ("Lectura anterior", "previous_reading"),
        ("Fecha de lectura actual", "current_reading_date"),
        ("Lectura actual", "current_reading"),
        ("Consumo (m³)", "consumption_m3"),
        ("Agua (€)", "water_amount"),
        ("Alcantarillado (€)", "sewerage_amount"),
        ("Canon de saneamiento (€)", "sanitation_amount"),
        ("IVA (%)", "vat_rate"),
        ("Base IVA (€)", "vat_base"),
        ("IVA (€)", "vat_amount"),
    ),
    "luz": (
        ("CUPS", "cups"),
        ("Contador", "meter_number"),
        ("Potencia contratada (kW)", "contracted_power_kw"),
        ("Potencia contratada valle (kW)", "contracted_power_valley_kw"),
        ("Consumo (kWh)", "consumption_kwh"),
        ("Consumo punta (kWh)", "consumption_peak_kwh"),
        ("Consumo llano (kWh)", "consumption_flat_kwh"),
        ("Consumo valle (kWh)", "consumption_valley_kwh"),
        ("Días facturados", "billing_days"),
        ("Tarifa", "tariff"),
        ("Peaje de acceso", "access_tariff"),
        ("Término fijo (€)", "fixed_term_amount"),
        ("Energía (€)", "energy_amount"),
        ("Otros conceptos (€)", "other_concepts_amount"),
        ("Impuesto eléctrico (€)", "electricity_tax_amount"),
        ("Base impuesto eléctrico (€)", "electricity_tax_base"),
        ("IVA (€)", "vat_amount"),
        ("Base IVA (€)", "vat_base"),
    ),
    "gas": (
        ("CUPS", "cups"),
        ("Contador", "meter_number"),
        ("Distribuidora", "distributor"),
        ("Tipo de tarifa", "tariff"),
        ("Tarifa de acceso", "access_tariff"),
        ("Fecha de lectura anterior", "previous_reading_date"),
        ("Lectura anterior (m³)", "previous_reading"),
        ("Fecha de lectura actual", "current_reading_date"),
        ("Lectura actual (m³)", "current_reading"),
        ("Consumo (m³)", "consumption_m3"),
        ("Factor de conversión (kWh/m³)", "conversion_factor_kwh_m3"),
        ("Consumo (kWh)", "consumption_kwh"),
        ("Término fijo (€)", "fixed_term_amount"),
        ("Energía (€)", "energy_amount"),
        ("Alquiler de contador (€)", "meter_rental_amount"),
        ("Impuesto hidrocarburos (€)", "hydrocarbon_tax_amount"),
        ("Importe ATR (€)", "atr_amount"),
        ("Otros cargos (€)", "charges_amount"),
        ("IVA (%)", "vat_rate"),
        ("Base IVA (€)", "vat_base"),
        ("IVA (€)", "vat_amount"),
    ),
}

REQUIRED_FIELDS = ("invoice_number", "issue_date", "total_amount")
SERVICE_REQUIRED_FIELDS = {
    "agua": ("meter_number", "consumption_m3"),
    "luz": ("cups", "consumption_kwh", "contracted_power_kw"),
    "gas": ("cups", "consumption_kwh"),
}
NUMERIC_FIELDS = {
    "tax_base", "tax_amount", "total_amount", "previous_reading", "current_reading",
    "consumption_m3", "consumption_kwh", "contracted_power_kw", "contracted_power_valley_kw",
    "meter_size_mm", "water_amount", "sewerage_amount", "sanitation_amount", "vat_rate", "vat_base", "vat_amount",
    "consumption_peak_kwh", "consumption_flat_kwh", "consumption_valley_kwh", "billing_days",
    "fixed_term_amount", "energy_amount", "other_concepts_amount", "electricity_tax_amount",
    "electricity_tax_base", "vat_amount", "vat_base", "conversion_factor_kwh_m3",
    "meter_rental_amount", "hydrocarbon_tax_amount", "atr_amount", "charges_amount",
}
MISSING_LABELS = {
    "invoice_number": "número de factura",
    "issue_date": "fecha de emisión",
    "total_amount": "total",
    "cups": "CUPS",
    "consumption_m3": "consumo en m³",
    "meter_number": "contador",
    "consumption_kwh": "consumo en kWh",
    "contracted_power_kw": "potencia contratada",
}


def _label_value(text: str, labels: str) -> str:
    match = re.search(
        rf"(?:{labels})\s*(?:n[.º°]?\s*)?[:#\-]?\s*([^\r\n|]{{2,100}})",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    value = re.sub(r"\s{2,}", " ", match.group(1)).strip(" \t:;-")
    return "" if value.casefold() in {"", "n/a", "no disponible"} else value


def _line_label_value(text: str, labels: str) -> str:
    match = re.search(rf"(?im)^\s*(?:{labels})\s*[:#\-]?\s*([^\r\n]+)", text)
    return re.sub(r"\s{2,}", " ", match.group(1)).strip(" \t:;-") if match else ""


def _water_component_amount(text: str, label: str) -> str:
    match = re.search(
        rf"(?im)^\s*{label}(?:\s*\(\d+\))?\s+(\d[\d. ]*,\d{{2}})",
        text,
    )
    return match.group(1).replace(" ", "") if match else ""


def _amount_after_label(text: str, labels: str) -> str:
    match = re.search(
        rf"(?:{labels})[^\d\r\n]{{0,40}}(\d[\d. ]*(?:,\d{{2}}|\.\d{{2}}))\s*(?:€|eur)?",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).replace(" ", "") if match else ""


def _amount_with_unit(text: str, labels: str, unit: str) -> str:
    match = re.search(
        rf"(?:{labels})[^\d\r\n]{{0,40}}(\d[\d. ]*(?:,\d{{1,3}}|\.\d{{1,3}})?)\s*{unit}\b",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).replace(" ", "") if match else ""


def _amounts_with_unit(text: str, labels: str, unit: str) -> list[str]:
    match = re.search(rf"(?:{labels})[^\r\n]*", text, flags=re.IGNORECASE)
    if not match:
        return []
    return [
        value.replace(" ", "")
        for value in re.findall(
            rf"(\d[\d. ]*(?:,\d{{1,3}}|\.\d{{1,3}})?)\s*{unit}\b",
            match.group(0),
            flags=re.IGNORECASE,
        )
    ]


def parse_invoice_text(text: str, invoice_type: str, provider: str) -> dict[str, str]:
    """Extract common Spanish utility invoice values using local text rules."""
    fields = {key: "" for _, key in (*COMMON_FIELDS, *SERVICE_FIELDS[invoice_type])}
    fields.update({"invoice_type": invoice_type, "provider": provider, "extraction_status": "Revisar", "missing_fields": ""})

    fields["holder"] = _line_label_value(text, r"titular(?:\s+del\s+contrato)?")
    if not fields["holder"]:
        fields["holder"] = _line_label_value(text, r"cliente")
    fields["tax_id"] = _label_value(text, r"(?:NIF|CIF|DNI|NIE)(?:\s*/\s*(?:NIF|CIF))?")
    address_label = r"direcci[oó\ufffd]n\s+(?:de(?:l)?\s+)?(?:suministro|punto\s+de\s+suministro)"
    address_match = re.search(
        rf"(?:{address_label})\s*[:#\-]?\s*([^\r\n|]{{2,100}})",
        text,
        flags=re.IGNORECASE,
    )
    if invoice_type == "agua" and not address_match:
        address_match = re.search(
            r"(?im)^\s*direcci[oó\ufffd]n\s*[:#\-]?\s*([^\r\n]*)",
            text,
        )
    fields["supply_address"] = address_match.group(1).strip(" \t:;-") if address_match else ""
    if address_match:
        following_line = text[address_match.end():].lstrip("\r\n").splitlines()
        if following_line:
            continuation = following_line[0].strip()
            next_field = re.match(
                r"(?i)^(?:cups\b|n[.º°\ufffd]?\s*(?:de\s+)?(?:factura|contrato|contador)\b|(?:fecha|total|periodo|factura|titular|nif|cif|potencia|consumo)\b)",
                continuation,
            )
            if continuation and not next_field and len(continuation) <= 80:
                fields["supply_address"] = f"{fields['supply_address']} {continuation}"
    fields["invoice_number"] = _label_value(
        text,
        r"n[uú\ufffd]mero\s+de\s+factura|n[.º°\ufffd]\s*(?:de\s+)?factura|factura\s+n[uú\ufffd]m(?:ero)?\.?|n[uú\ufffd]m(?:ero)?\.?\s+factura",
    )
    fields["contract_number"] = _label_value(
        text,
        r"(?:n[uú\ufffd]?m(?:ero)?\.?|n[.º°\ufffd])\s*(?:de\s+)?contrato",
    )
    fields["access_contract_number"] = _label_value(
        text,
        r"n[.º°\ufffd]\s*(?:de\s+)?contrato\s+de\s+acceso|n[uú\ufffd]mero\s+de\s+contrato\s+de\s+acceso",
    )

    date_pattern = r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})"
    for key, labels in (
        ("issue_date", r"fecha\s+(?:de\s+)?(?:emisi[oó\ufffd]n|factura)"),
        ("period_start", r"(?:periodo|per[ií\ufffd]odo)\s+(?:de\s+)?(?:consumo\s+)?desde"),
        ("period_end", r"(?:periodo|per[ií\ufffd]odo)\s+(?:de\s+)?(?:consumo\s+)?hasta"),
        ("due_date", r"(?:fecha\s+de\s+)?(?:vencimiento|cobro|cargo)|pagar\s+antes\s+de"),
    ):
        match = re.search(rf"(?:{labels})[^\d\r\n]{{0,24}}{date_pattern}", text, flags=re.IGNORECASE)
        if match:
            fields[key] = match.group(1)

    period_range = re.search(
        rf"(?:periodo|per[ií\ufffd]odo)\s+de\s+facturaci[oó\ufffd]n[^\d\r\n]{{0,24}}{date_pattern}\s*[-–a]\s*{date_pattern}",
        text,
        flags=re.IGNORECASE,
    )
    if period_range:
        fields["period_start"], fields["period_end"] = period_range.group(1), period_range.group(2)

    if invoice_type == "gas":
        gas_invoice = re.search(
            rf"fecha\s+y\s+n[.º°\ufffd]?\s*factura\s*:?\s*{date_pattern}\s*/\s*([A-Z0-9-]+)",
            text,
            flags=re.IGNORECASE,
        )
        if gas_invoice:
            fields["issue_date"], fields["invoice_number"] = gas_invoice.group(1), gas_invoice.group(2)
        gas_period = re.search(
            rf"periodo\s+de\s+facturaci[oó\ufffd]n\s*:?\s*(?:del\s+)?{date_pattern}\s+al\s+{date_pattern}",
            text,
            flags=re.IGNORECASE,
        )
        if gas_period:
            fields["period_start"], fields["period_end"] = gas_period.group(1), gas_period.group(2)

    fields["tax_base"] = _amount_after_label(text, r"base\s+imponible")
    fields["tax_amount"] = _amount_after_label(text, r"(?:total\s+)?impuestos|impuestos\s+totales|IVA")
    fields["total_amount"] = _amount_after_label(
        text,
        r"total\s+a\s+pagar|importe\s+total|total\s+factura|(?<![\w])total(?=\s*[:\-])",
    )

    cups = re.search(r"\b([A-Z]{2}\d{16,20}[A-Z0-9]{2,4})\b", text, flags=re.IGNORECASE)
    if cups and invoice_type in {"luz", "gas"}:
        fields["cups"] = cups.group(1).upper()

    if invoice_type == "agua":
        meter_row = re.search(
            r"(?im)^\s*(?P<meter>[A-Z0-9]{6,})\s+(?P<size>\d{1,3})\s+"
            r"(?P<previous_date>\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\s+"
            r"(?P<previous>\d[\d.,]*)\s+"
            r"(?P<current_date>\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\s+"
            r"(?P<current>\d[\d.,]*)\s+(?P<consumption>\d[\d.,]*)\s+(?:Real|Estimada|Estimado)\b",
            text,
        )
        if meter_row:
            values = meter_row.groupdict()
            fields["meter_number"] = values["meter"]
            fields["meter_size_mm"] = values["size"]
            fields["previous_reading_date"] = values["previous_date"]
            fields["previous_reading"] = values["previous"]
            fields["current_reading_date"] = values["current_date"]
            fields["current_reading"] = values["current"]
            fields["consumption_m3"] = values["consumption"]
            fields["period_start"] = fields["period_start"] or values["previous_date"]
            fields["period_end"] = fields["period_end"] or values["current_date"]
        else:
            fields["meter_number"] = _line_label_value(
                text,
                r"(?:n[uú\ufffd]?m(?:ero)?\.?\s+de\s+)?contador",
            )
            fields["previous_reading"] = _amount_after_label(text, r"lectura\s+anterior")
            fields["current_reading"] = _amount_after_label(text, r"lectura\s+actual|lectura\s+nueva")
        explicit_consumption = _amount_with_unit(text, r"consumo(?:\s+de\s+agua)?", r"m[³3]")
        if explicit_consumption:
            fields["consumption_m3"] = explicit_consumption

        fields["water_amount"] = _water_component_amount(text, r"AGUA")
        fields["sewerage_amount"] = _water_component_amount(text, r"ALCANTARILLADO")
        fields["sanitation_amount"] = _water_component_amount(text, r"CANON\s+(?:DE\s+)?SANEAMIENTO")
        vat_match = re.search(
            r"(?im)^\s*I\.?\s*V\.?\s*A\.?\s+al\s+(\d{1,2})\s*%\s*"
            r"BASE\s+IMPONIBLE\s*:?\s*(\d[\d. ]*,\d{2})\s+(\d[\d. ]*,\d{2})",
            text,
        )
        if vat_match:
            fields["vat_rate"], fields["vat_base"], fields["vat_amount"] = (
                value.replace(" ", "") for value in vat_match.groups()
            )
            fields["tax_base"] = fields["vat_base"]
            fields["tax_amount"] = fields["vat_amount"]
    elif invoice_type == "gas":
        meter_match = re.search(
            r"(?:n[.º°\ufffd]?\s*(?:de\s+)?contador|contador\s+n[.º°\ufffd]?)\s*:?\s*([A-Z0-9-]{5,})",
            text,
            flags=re.IGNORECASE,
        )
        fields["meter_number"] = meter_match.group(1) if meter_match else ""
        tariff_type = re.search(
            r"(?im)\btipo\s*:?\s*([^\r\n|]*?)(?=\s+n[.º°\ufffd]?\s*(?:de\s+)?contrato\b|\s*$)",
            text,
        )
        fields["tariff"] = tariff_type.group(1).strip(" \t:;-") if tariff_type else ""
        access_tariff = re.search(r"tarifa\s+de\s+acceso\s*:?\s*([A-Z0-9.]+)", text, flags=re.IGNORECASE)
        fields["access_tariff"] = access_tariff.group(1).upper() if access_tariff else ""
        distributor = re.search(
            r"(?im)^\s*distribuidora\s*:?\s*([^\r\n|]{2,100})",
            text,
        )
        fields["distributor"] = distributor.group(1).strip(" \t:;-") if distributor else ""

        reading_pattern = (
            r"lectura\s+(actual|anterior)\s+(?:real|estimada)?\s*"
            r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\s+(\d[\d., ]*)\s*m[³3]"
        )
        for reading in re.finditer(reading_pattern, text, flags=re.IGNORECASE):
            kind, reading_date, value = reading.groups()
            prefix = "current" if kind.casefold() == "actual" else "previous"
            fields[f"{prefix}_reading_date"] = reading_date
            fields[f"{prefix}_reading"] = value.strip().replace(" ", "")

        fields["consumption_m3"] = _amount_with_unit(text, r"consumo\s+en\s+m[³3]", r"m[³3]")
        conversion = re.search(
            r"factor\s+de\s+conversi[oó\ufffd]n[^\d\r\n]{0,24}(\d[\d., ]*)\s*kWh\s*/\s*m[³3]",
            text,
            flags=re.IGNORECASE,
        )
        fields["conversion_factor_kwh_m3"] = conversion.group(1).replace(" ", "") if conversion else ""
        fields["consumption_kwh"] = _amount_with_unit(text, r"consumo\s+en\s+kWh", r"kWh")
        fixed_line = re.search(r"(?im)^\s*t[eé\ufffd]rmino\s+fijo[^\r\n]*", text)
        if fixed_line:
            fixed_amounts = re.findall(r"\d[\d. ]*,\d{2}(?=\s*(?:€|$))", fixed_line.group(0))
            if fixed_amounts:
                fields["fixed_term_amount"] = fixed_amounts[-1].replace(" ", "")
        energy_line = re.search(
            r"(?im)^\s*t[eé\ufffd]rmino\s+energ[ií\ufffd]a[^\r\n]*",
            text,
        )
        if energy_line:
            energy_amounts = re.findall(r"\d[\d. ]*,\d{2}(?=\s*(?:€|$))", energy_line.group(0))
            if energy_amounts:
                fields["energy_amount"] = energy_amounts[-1].replace(" ", "")
        fields["meter_rental_amount"] = _amount_after_label(text, r"alquiler\s+de\s+contador")
        hydrocarbon_line = re.search(
            r"(?im)^\s*impuesto\s+(?:esp\.?\s*)?hidrocarburos[^\r\n]*",
            text,
        )
        if hydrocarbon_line:
            amounts = re.findall(r"\d[\d. ]*,\d{2}(?=\s*(?:€|$))", hydrocarbon_line.group(0))
            if amounts:
                fields["hydrocarbon_tax_amount"] = amounts[-1].replace(" ", "")
        for key, label in (("atr_amount", r"importe\s+ATR"), ("charges_amount", r"importe\s+cargos")):
            amount = re.search(
                rf"{label}\s*:?\s*(\d[\d. ]*,\d{{2}})(?=\s*€|\s+importe|\s*$)",
                text,
                flags=re.IGNORECASE,
            )
            if amount:
                fields[key] = amount.group(1).replace(" ", "")

        gas_vat = re.search(
            r"I\.?\s*V\.?\s*A\.?\s*\(\s*(\d{1,2})\s*%\s*\)[^\r\n]*?"
            r"(\d[\d. ]*,\d{2})\s+(\d[\d. ]*,\d{2})",
            text,
            flags=re.IGNORECASE,
        )
        if gas_vat:
            fields["vat_rate"], fields["vat_base"], fields["vat_amount"] = (
                value.replace(" ", "") for value in gas_vat.groups()
            )
            fields["tax_base"] = fields["vat_base"]
            fields["tax_amount"] = fields["vat_amount"]
    else:
        fields["consumption_kwh"] = _amount_with_unit(text, r"(?:consumo|energ[ií]a)(?:\s+total)?", r"kWh")
        fields["tariff"] = _label_value(text, r"tarifa|peaje\s+de\s+acceso")
        if invoice_type == "luz":
            fields["contracted_power_kw"] = _amount_with_unit(text, r"potencia\s+contratada", r"kW")
            power_values = _amounts_with_unit(text, r"potencia\s+contratada", r"kW")
            if len(power_values) > 1:
                fields["contracted_power_valley_kw"] = power_values[1]
            fields["meter_number"] = _label_value(
                text,
                r"(?:n[uú\ufffd]mero\s+de\s+)?contador|n[.º°\ufffd]\s*(?:de\s+)?contador|n[uú\ufffd]m\.\s*contador",
            )
            fields["billing_days"] = _amount_with_unit(text, r"d[ií\ufffd]as\s+facturados", r"d[ií\ufffd]as?")
            consumption_period = re.search(
                r"periodo\s+de\s+consumo\s+punta\s+llano\s+valle(.{0,600})",
                text,
                flags=re.IGNORECASE | re.DOTALL,
            )
            if consumption_period:
                period_values = re.findall(
                    r"(\d[\d. ]*(?:,\d{1,3}|\.\d{1,3})?)\s*kWh\b",
                    consumption_period.group(1),
                    flags=re.IGNORECASE,
                )[:3]
                period_keys = ("consumption_peak_kwh", "consumption_flat_kwh", "consumption_valley_kwh")
                for key, value in zip(period_keys, period_values):
                    fields[key] = value.replace(" ", "")
            access_tariff = re.search(r"\b\d\.\d(?:TD|RL)\b", text, flags=re.IGNORECASE)
            fields["access_tariff"] = access_tariff.group(0).upper() if access_tariff else ""
            fields["fixed_term_amount"] = _amount_after_label(text, r"t[eé\ufffd]rmino\s+fijo")
            fields["energy_amount"] = _amount_after_label(text, r"energ[ií\ufffd]a")
            fields["other_concepts_amount"] = _amount_after_label(text, r"otros\s+conceptos")
            fields["electricity_tax_amount"] = _amount_after_label(text, r"impuesto\s+el[eé\ufffd]ctrico")
            electricity_tax_match = re.search(
                r"impuesto\s+el[eé\ufffd]ctrico[^\d\r\n]{0,24}\d[\d. ]*(?:,\d{2}|\.\d{2})\s*€?\s+(\d[\d. ]*(?:,\d{2}|\.\d{2}))",
                text,
                flags=re.IGNORECASE,
            )
            if electricity_tax_match:
                fields["electricity_tax_base"] = electricity_tax_match.group(1).replace(" ", "")
            vat_match = re.search(
                r"IVA\s*\([^\r\n)]*%\)[^\d\r\n]{0,24}(\d[\d. ]*(?:,\d{2}|\.\d{2}))\s+(\d[\d. ]*(?:,\d{2}|\.\d{2}))",
                text,
                flags=re.IGNORECASE,
            )
            if vat_match:
                fields["vat_base"] = vat_match.group(1).replace(" ", "")
                fields["vat_amount"] = vat_match.group(2).replace(" ", "")
        else:
            fields["meter_number"] = _label_value(
                text,
                r"(?:n[uú\ufffd]mero\s+de\s+)?contador|n[.º°\ufffd]\s*(?:de\s+)?contador|n[uú\ufffd]m\.\s*contador",
            )

    required_fields = (*REQUIRED_FIELDS, *SERVICE_REQUIRED_FIELDS[invoice_type])
    missing = [key for key in required_fields if not fields.get(key)]
    fields["missing_fields"] = ", ".join(MISSING_LABELS.get(key, key) for key in missing)
    fields["extraction_status"] = "Correcta" if not missing else "Revisar"
    return fields


def extract_invoice_fields(pdf_path: str | Path, invoice_type: str, provider: str) -> dict[str, str]:
    """Use embedded PDF text first, then local Spanish/English OCR if needed."""
    started = perf_counter()
    ocr_used = False
    fields: dict[str, str] = {}
    try:
        with fitz.open(pdf_path) as pdf:
            page_texts = [page.get_text("text") for page in pdf]
            text = "\n".join(page_texts)
            fields = parse_invoice_text(text, invoice_type, provider)
            if fields["extraction_status"] == "Correcta":
                return fields

            try:
                ocr_used = True
                ocr_text = []
                with fitz.open(pdf_path) as pdf:
                    for page in pdf:
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                        ocr_text.append(pytesseract.image_to_string(image, lang="spa+eng", config="--psm 6"))
                ocr_fields = parse_invoice_text("\n".join([text, *ocr_text]), invoice_type, provider)
                if len(ocr_fields["missing_fields"].split(", ")) < len(fields["missing_fields"].split(", ")):
                    fields = ocr_fields
            except Exception as exc:  # OCR is a fallback; retain any embedded text already found.
                fields["ocr_note"] = f"OCR no disponible: {exc}"
            return fields
    except Exception as exc:
        fields = parse_invoice_text("", invoice_type, provider)
        fields["ocr_note"] = f"No se pudo leer el PDF: {exc}"
        return fields
    finally:
        logger.info(
            "Invoice extraction timing provider=%s service=%s elapsed_ms=%.0f ocr_used=%s status=%s",
            provider,
            invoice_type,
            (perf_counter() - started) * 1000,
            ocr_used,
            fields.get("extraction_status", "unknown"),
        )


def build_invoice_workbook(documents: list, invoice_type: str, destination: str | Path) -> Path:
    started = perf_counter()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = invoice_type.capitalize()
    fields = [*COMMON_FIELDS, *SERVICE_FIELDS[invoice_type], ("Estado de extracción", "extraction_status"), ("Campos pendientes", "missing_fields")]
    sheet.append([label for label, _ in fields])
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for document in documents:
        extracted = extract_invoice_fields(document.file_path, invoice_type, document.provider)
        row = []
        for _, key in fields:
            value = extracted.get(key, "")
            if value and key in NUMERIC_FIELDS:
                normalized = value.replace(" ", "")
                if "," in normalized:
                    normalized = normalized.replace(".", "").replace(",", ".")
                elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", normalized):
                    normalized = normalized.replace(".", "")
                try:
                    value = float(Decimal(normalized))
                except InvalidOperation:
                    pass
            row.append(value)
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
        sheet.column_dimensions[letter].width = width
    headers = {cell.value: cell.column for cell in sheet[1]}
    for label, key in fields:
        if key in NUMERIC_FIELDS:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, headers[label]).number_format = "#,##0.00"
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    logger.info(
        "Invoice workbook timing service=%s documents=%d elapsed_ms=%.0f",
        invoice_type,
        len(documents),
        (perf_counter() - started) * 1000,
    )
    return destination
